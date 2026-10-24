"""SQL wall times survive export checksums, privacy, and artifact rendering."""

import csv
import json
from datetime import UTC, datetime
from pathlib import Path

import duckdb
import pytest
from openpyxl import load_workbook

from moneybin.database import Database
from moneybin.exports.models import RedactionMode
from moneybin.exports.renderers import render_csv, render_parquet, render_xlsx
from moneybin.exports.service import ExportService
from moneybin.privacy.taxonomy import DataClass
from moneybin.reports._framework.catalog import ReportCatalog
from moneybin.reports._framework.contract import OutputColumn, ReportQuery, ReportSpec
from moneybin.services.user_reports_service import UserReportsService
from moneybin.tables import TableRef
from tests.moneybin.db_helpers import create_core_tables_raw
from tests.moneybin.test_reports._metadata import TEST_SEMANTICS


def _service(
    db: Database,
    query: str,
    *,
    timestamp_class: DataClass = DataClass.TIMESTAMP_OBSERVABILITY,
) -> ExportService:
    def runner(database: Database) -> ReportQuery:
        return ReportQuery(query, [])

    classes = {
        "wall": timestamp_class,
        "instant": timestamp_class,
        "nested": timestamp_class,
        "account_number": DataClass.ACCOUNT_IDENTIFIER,
    }
    spec = ReportSpec(
        report_id="test:wall_time",
        name="wall_time",
        description="Synthetic computed timestamp report.",
        view=TableRef("reports", "wall_time"),
        runner=runner,
        columns=tuple(OutputColumn(name, name, dc) for name, dc in classes.items()),
        semantics=TEST_SEMANTICS,
        classes=classes,
        examples=(),
    )
    return ExportService(db, report_catalog=ReportCatalog((spec,)))


@pytest.mark.parametrize("mode", ["redacted", "unredacted"])
def test_report_exports_preserve_computed_wall_times(
    db: Database, tmp_path: Path, mode: RedactionMode
) -> None:
    service = _service(
        db,
        """
        SELECT date_trunc('month', DATE '2026-10-21') AS wall,
               TIMESTAMPTZ '2026-10-01 09:00:00.123456-04:00' AS instant,
               {'times': [TIMESTAMP '2026-10-01 09:00:00.123456']} AS nested,
               'acct_11112222' AS account_number
        """,
    )

    snapshot = service.prepare_report(
        profile="test", report_id="test:wall_time", redaction_mode=mode
    )

    account = "****2222" if mode == "redacted" else "acct_11112222"
    wall = datetime(2026, 10, 1)
    instant = datetime(2026, 10, 1, 13, 0, 0, 123456, tzinfo=UTC)
    nested = {"times": [datetime(2026, 10, 1, 9, 0, 0, 123456)]}
    assert snapshot.tables[0].rows == ((wall, instant, nested, account),)
    assert [column.duckdb_type for column in snapshot.tables[0].columns] == [
        "TIMESTAMP",
        "TIMESTAMP WITH TIME ZONE",
        "STRUCT(times TIMESTAMP[])",
        "VARCHAR",
    ]
    expected_text = [
        "2026-10-01T00:00:00.000000",
        "2026-10-01T13:00:00.123456Z",
        '{"times":["2026-10-01T09:00:00.123456"]}',
        account,
    ]
    csv_artifact = render_csv(snapshot, tmp_path / "csv")
    with csv_artifact.table_files["test:wall_time"].open(newline="") as stream:
        assert list(csv.reader(stream))[1] == expected_text
    workbook = load_workbook(render_xlsx(snapshot, tmp_path / "xlsx").path)
    try:
        row = next(workbook.worksheets[0].iter_rows(min_row=2, values_only=True))
        assert list(row) == expected_text
    finally:
        workbook.close()
    parquet = render_parquet(snapshot, tmp_path / "parquet")
    assert duckdb.read_parquet(
        str(parquet.table_files["test:wall_time"])
    ).fetchall() == [(wall, instant, nested, account)]
    assert (
        json.loads((csv_artifact.path / "manifest.json").read_text())["tables"][0][
            "checksum_sha256"
        ]
        == snapshot.tables[0].checksum_sha256
    )


@pytest.mark.parametrize("timestamp_type", ["TIMESTAMP", "TIMESTAMPTZ"])
@pytest.mark.parametrize("value", ["infinity", "-infinity"])
def test_report_exports_keep_refusing_infinite_timestamps(
    db: Database, timestamp_type: str, value: str
) -> None:
    service = _service(
        db,
        f"SELECT {timestamp_type} '{value}' AS wall",
    )

    with pytest.raises(ValueError, match="explicit UTC offset"):
        service.prepare_report(profile="test", report_id="test:wall_time")


def test_saved_report_exports_computed_wall_times(db: Database, tmp_path: Path) -> None:
    create_core_tables_raw(db.conn)
    db.execute(
        """
        INSERT INTO core.fct_transactions (transaction_id, transaction_date)
        VALUES ('txn_wall_time', DATE '2026-10-21')
        """
    )
    report_id = (
        UserReportsService(db)
        .create(
            name="computed_times",
            query_sql="""
        SELECT date_trunc('month', transaction_date) AS wall,
               {'times': [date_trunc('month', transaction_date)]} AS nested,
               TIMESTAMPTZ '2026-10-01 09:00:00.123456-04:00' AS instant
        FROM core.fct_transactions
        """,
            actor="cli",
        )
        .report_id
    )

    snapshot = ExportService(db).prepare_report(profile="test", report_id=report_id)
    expected = [
        "2026-10-01T00:00:00.000000",
        '{"times":["2026-10-01T00:00:00.000000"]}',
        "2026-10-01T13:00:00.123456Z",
    ]
    artifact = render_csv(snapshot, tmp_path / "csv")
    with artifact.table_files[report_id].open(newline="") as stream:
        assert list(csv.reader(stream))[1] == expected


def test_floored_report_export_formats_timestamps_after_privacy(
    db: Database, tmp_path: Path
) -> None:
    service = _service(
        db,
        """
        SELECT TIMESTAMP '2026-10-01 00:00:00' AS wall,
               {'times': [TIMESTAMP '2026-10-01 09:00:00.123456'],
                'routing': '021000021'} AS nested,
               TIMESTAMPTZ '2026-10-01 09:00:00.123456-04:00' AS instant
        """,
        timestamp_class=DataClass.FLOORED,
    )

    snapshot = service.prepare_report(profile="test", report_id="test:wall_time")
    expected = [
        "2026-10-01T00:00:00.000000",
        "{'times': ['2026-10-01T09:00:00.123456'], 'routing': '****...0021'}",
        "2026-10-01T13:00:00.123456Z",
    ]
    artifact = render_csv(snapshot, tmp_path / "csv")
    with artifact.table_files["test:wall_time"].open(newline="") as stream:
        assert list(csv.reader(stream))[1] == expected
    parquet = render_parquet(snapshot, tmp_path / "parquet")
    assert duckdb.read_parquet(
        str(parquet.table_files["test:wall_time"])
    ).fetchall() == [tuple(expected)]


@pytest.mark.parametrize(
    ("expression", "expected_key"),
    [
        ("TIMESTAMP '2026-10-01 09:00:00.123456'", "2026-10-01T09:00:00.123456"),
        (
            "TIMESTAMPTZ '2026-10-01 09:00:00.123456-04:00'",
            "2026-10-01T13:00:00.123456Z",
        ),
    ],
)
@pytest.mark.parametrize(
    "data_class", [DataClass.TIMESTAMP_OBSERVABILITY, DataClass.FLOORED]
)
def test_report_exports_format_timestamp_map_keys(
    db: Database,
    tmp_path: Path,
    expression: str,
    expected_key: str,
    data_class: DataClass,
) -> None:
    service = _service(
        db,
        f"SELECT map([{expression}], [7]) AS wall",
        timestamp_class=data_class,
    )
    snapshot = service.prepare_report(profile="test", report_id="test:wall_time")
    expected = (
        f"{{'{expected_key}': 7}}"
        if data_class is DataClass.FLOORED
        else f'{{"{expected_key}":7}}'
    )

    artifact = render_csv(snapshot, tmp_path / "csv")
    with artifact.table_files["test:wall_time"].open(newline="") as stream:
        assert list(csv.reader(stream))[1] == [expected]
    if data_class is DataClass.TIMESTAMP_OBSERVABILITY:
        assert snapshot.tables[0].rows == (
            ({datetime.fromisoformat(expected_key): 7},),
        )


@pytest.mark.parametrize("timestamp_type", ["TIMESTAMP", "TIMESTAMPTZ"])
@pytest.mark.parametrize("value", ["infinity", "-infinity"])
def test_report_exports_refuse_infinite_timestamp_map_keys(
    db: Database, timestamp_type: str, value: str
) -> None:
    service = _service(
        db,
        f"SELECT map([{timestamp_type} '{value}'], [7]) AS wall",
    )

    with pytest.raises(ValueError, match="explicit UTC offset"):
        service.prepare_report(profile="test", report_id="test:wall_time")
